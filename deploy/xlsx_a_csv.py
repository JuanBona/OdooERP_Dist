"""Convierte los Excel del cliente a CSV UTF-8 para deploy/carga_inicial.py.

    python deploy/xlsx_a_csv.py <Clientes.xlsx> <ListaPrecios.xlsx> <carpeta_salida>

Genera clientes.csv y productos.csv. Los CSV contienen datos de clientes:
NO commitearlos (el repo es público).
"""
import csv
import os
import sys

import openpyxl

clientes_xlsx, precios_xlsx, out = sys.argv[1:4]
os.makedirs(out, exist_ok=True)


def limpio(v):
    v = "" if v is None else str(v).strip()
    return "" if v in ("(0)", "xx", "XX") else v


with open(os.path.join(out, "clientes.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["codigo", "nombre", "domicilio", "cuit", "telefono", "email"])
    ws = openpyxl.load_workbook(clientes_xlsx, read_only=True, data_only=True).worksheets[0]
    n = 0
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] is None or not limpio(r[1]):
            continue
        w.writerow([int(r[0]), limpio(r[1]), limpio(r[2]), limpio(r[3]), limpio(r[4]), limpio(r[5])])
        n += 1
print("clientes:", n)

with open(os.path.join(out, "productos.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["categoria", "nombre", "precio"])
    n = 0
    for ws in openpyxl.load_workbook(precios_xlsx, read_only=True, data_only=True).worksheets:
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not limpio(r[0]) or r[2] is None:
                continue
            w.writerow([ws.title.strip(), limpio(r[0]), r[2]])
            n += 1
print("productos:", n)
